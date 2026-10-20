"""Unit tests for AI Training Studio providers: document extraction, LLM clients,

embeddings, and video rendering.
"""

import io
import wave

import docx
import httpx
import pptx
import pymupdf
import pytest

from app.modules.training_studio import embeddings, llm_client, parser, renderer_client
from app.modules.training_studio.schemas import (
    CourseDraft,
    LessonDraft,
    ModuleDraft,
    StoryboardScene,
)


def test_docx_extraction_with_python_docx():
    """Verify python-docx parses structured paragraphs and tables."""
    doc = docx.Document()
    doc.add_heading("Emergency Response Procedure", level=1)
    doc.add_paragraph("All personnel must assemble at designated evacuation zones.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Zone"
    table.rows[0].cells[1].text = "Coordinator"
    table.rows[1].cells[0].text = "North Gate"
    table.rows[1].cells[1].text = "Security Lead"

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    extracted = parser.extract_text_from_bytes(file_bytes, "emergency_procedure.docx")
    assert "Emergency Response Procedure" in extracted
    assert "evacuation zones" in extracted
    assert "North Gate" in extracted


def test_pptx_extraction_with_python_pptx():
    """Verify python-pptx extracts slide text frames."""
    prs = pptx.Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Quarterly Compliance Briefing"
    subtitle = slide.placeholders[1]
    subtitle.text = "Mandatory briefing for all financial analysts."

    buf = io.BytesIO()
    prs.save(buf)
    file_bytes = buf.getvalue()

    extracted = parser.extract_text_from_bytes(file_bytes, "briefing.pptx")
    assert "Quarterly Compliance Briefing" in extracted
    assert "Mandatory briefing" in extracted


def test_pdf_extraction_with_pymupdf():
    """Verify PyMuPDF extracts text streams cleanly."""
    doc = pymupdf.open()
    page = doc.new_page()
    text = "Anti-Money Laundering Safeguards\nTransaction threshold is $10,000."
    page.insert_text(pymupdf.Point(50, 72), text)
    pdf_bytes = doc.tobytes()

    extracted = parser.extract_text_from_bytes(pdf_bytes, "aml_policy.pdf")
    assert "Anti-Money Laundering Safeguards" in extracted
    assert "10,000" in extracted


def test_embedding_provider_deterministic_fallback():
    """Verify Ollama fallback produces normalized 384-d vectors."""
    provider = embeddings.get_embedding_provider()
    vec = provider.embed_text("Information Security Policy")

    assert isinstance(vec, list)
    assert len(vec) == 384
    # Check L2 normalization (sum of squares is ~1.0)
    norm = sum(v * v for v in vec)
    assert 0.99 <= norm <= 1.01


def test_llm_client_safe_fallback_on_invalid_json():
    """Verify JSON parser strips code fences and returns None on malformed output."""
    assert llm_client._parse_json('```json\n{"valid": true}\n```') == {"valid": True}
    assert llm_client._parse_json("Not a json string") is None


_SCENES = [
    StoryboardScene(
        scene_number=1,
        title="Intro",
        narration="Welcome to our compliance review.",
        visual_prompt="Welcome title card",
        duration_sec=3,
    ),
]


def _course_draft() -> CourseDraft:
    return CourseDraft(
        title="Phishing Defense",
        short_description="Short",
        description="Long description of the course content.",
        category="Security",
        difficulty_level="beginner",
        learning_objectives=["Spot phishing", "Report it", "Stay calm"],
        key_takeaways=["Never share credentials", "Report suspicious mail"],
        tags=["security"],
        is_mandatory=True,
        validity_days=365,
        modules=[
            ModuleDraft(
                title="Module 1: Signals",
                description="Signals",
                lessons=[
                    LessonDraft(
                        title="Urgency cues",
                        description="d",
                        content_type="text",
                        content="Text",
                        estimated_minutes=3,
                        is_required=True,
                    )
                ],
            )
        ],
    )


def test_render_scenes_map_objectives_lessons_and_takeaways():
    course = _course_draft()
    storyboard = [
        StoryboardScene(
            scene_number=1, title="Welcome to X", narration="n1", visual_prompt="v", duration_sec=6
        ),
        StoryboardScene(
            scene_number=2,
            title="Module 1: Signals",
            narration="n2",
            visual_prompt="v",
            duration_sec=10,
        ),
        StoryboardScene(
            scene_number=3,
            title="Summary & Knowledge Check",
            narration="n3",
            visual_prompt="v",
            duration_sec=6,
        ),
    ]
    scenes = renderer_client.build_render_scenes(course, storyboard)
    assert scenes[0]["bullets"] == ["Spot phishing", "Report it", "Stay calm"]
    assert scenes[1]["bullets"] == ["Urgency cues"]
    assert scenes[2]["bullets"] == ["Never share credentials", "Report suspicious mail"]


def test_narration_track_is_padded_to_scene_lengths():
    rate = 24000
    clip_a = (b"\x01\x00" * rate, rate)  # 1.0 s of audio
    clip_b = (b"\x02\x00" * (rate // 2), rate)  # 0.5 s of audio
    wav_bytes = renderer_client._build_narration_track([clip_a, clip_b], [2.0, 1.0])
    with wave.open(io.BytesIO(wav_bytes), "rb") as wf:
        assert wf.getframerate() == rate
        assert wf.getnframes() == 3 * rate  # 2.0 s + 1.0 s, each padded to its scene


def test_render_course_video_sends_scene_timings_and_returns_asset(tmp_path, monkeypatch):
    from app.core.config import get_settings

    monkeypatch.setattr(get_settings(), "gemini_api_key", None)
    captured: dict = {}

    def fake_post(url, json, timeout):  # noqa: A002
        captured["url"] = url
        captured["payload"] = json
        (tmp_path / f"video_{json['job_id']}.mp4").write_bytes(b"\x00" * 4096)
        return httpx.Response(200, json={"ok": True})

    monkeypatch.setattr(renderer_client.httpx, "post", fake_post)
    course = _course_draft()
    url, total = renderer_client.render_course_video(
        job_id="job-abc",
        course=course,
        storyboard=_SCENES,
        storage_dir=tmp_path,
    )
    assert url == "/api/v1/studio/video-asset/video_job-abc.mp4"
    assert total == 3
    assert captured["url"].endswith("/render")
    assert captured["payload"]["narration_wav_base64"] is None  # no key: silent, explicit
    assert captured["payload"]["scenes"][0]["duration_sec"] == 3.0


def test_render_course_video_surfaces_renderer_failure(tmp_path, monkeypatch):
    from app.core.config import get_settings

    monkeypatch.setattr(get_settings(), "gemini_api_key", None)

    def fake_post(url, json, timeout):  # noqa: A002
        return httpx.Response(500, json={"ok": False, "error": "chromium crashed"})

    monkeypatch.setattr(renderer_client.httpx, "post", fake_post)
    with pytest.raises(renderer_client.RenderError, match="chromium crashed"):
        renderer_client.render_course_video(
            job_id="job-fail",
            course=_course_draft(),
            storyboard=_SCENES,
            storage_dir=tmp_path,
        )
